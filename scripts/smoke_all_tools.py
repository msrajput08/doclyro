from pathlib import Path
import sys, tempfile,io,json,traceback
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app.processor import process, OPS
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from PIL import Image
import fitz
import datetime
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import rsa
from cryptography.hazmat.primitives.serialization import pkcs12
from cryptography.x509.oid import NameOID
from app.digital_signing import sign_pdf

with tempfile.TemporaryDirectory() as td:
 root=Path(td);d=fitz.open();p=d.new_page();p.insert_text((50,80),'Email: hello@example.com TEST PDF document.',fontsize=15);p.insert_text((50,130),'Name: Hello Example',fontsize=12);d.save(root/'a.pdf');d.close()
 doc=Document();doc.add_heading('Office Word Test',0);doc.add_paragraph('Hello Word to PDF');doc.save(root/'a.docx')
 prs=Presentation();sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.shapes.add_textbox(900000,900000,8000000,1000000).text='PowerPoint to PDF';prs.save(root/'a.pptx')
 wb=Workbook();ws=wb.active;ws.append(['Name','Score']);ws.append(['Hello',10]);wb.save(root/'a.xlsx')
 im=Image.new('RGB',(300,500),'white');im.save(root/'a.png')
 (root/'a.html').write_text('<h1>HTML to PDF</h1><p>HTML is working</p>')
 key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
 name=x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,'PDFCraft smoke signer')])
 today=datetime.datetime.now(datetime.timezone.utc)
 cert=(x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key())
       .serial_number(x509.random_serial_number()).not_valid_before(today-datetime.timedelta(days=1))
       .not_valid_after(today+datetime.timedelta(days=30)).sign(key,hashes.SHA256()))
 (root/'key.p12').write_bytes(pkcs12.serialize_key_and_certificates(
        b'smoke',key,cert,None,serialization.BestAvailableEncryption(b'password')))
 sign_pdf(root/'a.pdf',root/'key.p12','password',root/'signed.pdf')
 inputs={name:[root/'a.pdf'] for name in OPS}
 inputs['merge']=[root/'a.pdf',root/'a.pdf'];inputs['compare']=[root/'a.pdf',root/'a.pdf']
 inputs['sign-certificate']=[root/'a.pdf',root/'key.p12'];inputs['verify-signatures']=[root/'signed.pdf']
 for name,extension in [('jpg-to-pdf','.png'),('scan','.png'),('word-to-pdf','.docx'),('powerpoint-to-pdf','.pptx'),('excel-to-pdf','.xlsx'),('html-to-pdf','.html')]:inputs[name]=[root/('a'+extension)]
 opts={'split':{'mode':'each'},'remove-pages':{'pages':'1'},'extract-pages':{'pages':'1'},
 'organize':{'pages':'1'},'edit':{'text':'added'},'forms':{'mode':'create'},'protect':{'new_password':'myPass123!'},
 'sign':{'text':'Test Signer'},'sign-certificate':{'certificate_password':'password'},'redact':{'text':'Hello'},'translate':{'language':'de'},
 'workflow':{'steps':[{'tool':'rotate','options':{'angle':90}}]},'ocr':{'language':'eng'},
 'pdf-to-powerpoint':{'editable':'true'},'html-to-pdf':{},'pdf-to-pdfa':{}}
 failures={};success=[];skipped=[]
 for op in OPS:
  if op=='translate':skipped.append((op,'requires configured translator'));continue
  if op=='remove-pages':
   # single-page removal intentionally invalid; create a two-page input
   inputs[op]=[root/'two.pdf'];d=fitz.open();d.insert_pdf(fitz.open(root/'a.pdf'));d.insert_pdf(fitz.open(root/'a.pdf'));d.save(root/'two.pdf');d.close()
  try:
   wd=root/op;wd.mkdir();out=process(op,inputs[op],opts.get(op,{}),wd)
   if not out.exists() or out.stat().st_size<10:raise RuntimeError('No substantial output')
   if out.suffix.lower()=='.pdf':
    with fitz.open(out) as pdf:
     if not len(pdf):raise RuntimeError('Empty PDF')
   success.append((op,out.suffix.lower(),out.stat().st_size))
  except Exception as e:failures[op]=str(e)
 print('SUCCESS',len(success),'FAILED',len(failures),'SKIPPED',len(skipped))
 for op,ext,size in success:print('PASS',op,ext,size)
 for op,error in failures.items():print('FAIL',op,error[:400])
 for op,reason in skipped:print('SKIP',op,reason)

if failures:
 raise SystemExit('Tool smoke test failed')
