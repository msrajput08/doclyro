import sys
import tempfile
import unittest
from pathlib import Path
import fitz
from pptx import Presentation
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app.processor import process, ProcessingError

class ExtrasTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.wd=Path(self.temp.name)
        self.input=self.wd/'sample.pdf'
        with fitz.open() as pdf:
            p=pdf.new_page();p.insert_text((45,80),'PowerPoint editable text')
            pdf.save(str(self.input))
    def tearDown(self):self.temp.cleanup()
    def convert(self,tool,options=None):return process(tool,[self.input],options or {},self.wd)
    def test_pdf_to_pptx_has_real_text_boxes(self):
        out=self.convert('pdf-to-powerpoint',{'editable':'true'})
        prs=Presentation(out)
        self.assertEqual(len(prs.slides),1)
        self.assertTrue(any(shape.has_text_frame and 'PowerPoint editable text' in shape.text
                            for shape in prs.slides[0].shapes))
        self.assertTrue(any(shape.shape_type==13 for shape in prs.slides[0].shapes))
    def test_pdf_to_pptx_snapshot(self):
        out=self.convert('pdf-to-powerpoint',{'editable':'false'})
        prs=Presentation(out)
        self.assertEqual(len(prs.slides[0].shapes),1)
    def test_summarizer_extracts_when_unconfigured(self):
        out=self.convert('summarize',{'mode':'extractive'})
        self.assertIn('EXTRACTIVE',out.read_text())
    def test_forms_listbox(self):
        out=self.convert('forms',{'field_type':'field-listbox','field':'Country','choices':'IN,DE,US'})
        with fitz.open(out) as d:
            field=list(d[0].widgets())[0]
            self.assertEqual(field.field_type,fitz.PDF_WIDGET_TYPE_LISTBOX)
    def test_sticky_note_and_https_link(self):
        out=self.convert('edit',{'elements':[
            {'page':1,'type':'sticky-note','x':100,'y':90,'text':'Check date'},
            {'page':1,'type':'link','x':50,'y':120,'w':150,'h':25,'url':'https://example.com'}]})
        with fitz.open(out) as d:
            self.assertEqual(len(list(d[0].annots())),1)
            self.assertEqual(len(d[0].get_links()),1)
    def test_reject_non_https_link(self):
        with self.assertRaises(ProcessingError):self.convert('edit',{'elements':[
            {'page':1,'type':'link','x':50,'y':120,'w':150,'h':25,'url':'javascript:alert(1)'}]})

if __name__=='__main__':unittest.main()
