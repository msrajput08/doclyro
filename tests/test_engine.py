import io
import os
import tempfile
import unittest
from pathlib import Path
import fitz
from PIL import Image
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app.processor import process, pages, ProcessingError

class TestToolkit(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.wd=Path(self.tmp.name)
        self.pdf=self.wd/'sample.pdf';c=canvas.Canvas(str(self.pdf));c.drawString(50,700,'Hello PDFCraft World');c.showPage();c.drawString(60,720,'Secret Value 12345');c.save()
        self.second=self.wd/'second.pdf'; d=fitz.open();d.new_page().insert_text((72,72),'Appendix Document');d.save(str(self.second));d.close()
        self.image=self.wd/'scan.png';Image.new('RGB',(260,170),'white').save(self.image)
    def tearDown(self):self.tmp.cleanup()
    def go(self,slug,files=None,options=None):
        result=process(slug,files if files is not None else [self.pdf],options or {},self.wd)
        self.assertTrue(result.exists(),(slug,result))
        self.assertGreater(result.stat().st_size,0,slug)
        return result
    def assert_pages(self,path,n):
        with fitz.open(str(path)) as d:self.assertEqual(len(d),n)
    def test_page_parser(self):
        self.assertEqual(pages('1,3-4',4),[0,2,3])
        with self.assertRaises(ProcessingError):pages('4-1',4)
    def test_organize_tools(self):
        self.assert_pages(self.go('merge',[self.pdf,self.second]),3)
        self.assert_pages(self.go('remove-pages',options={'pages':'2'}),1)
        self.assert_pages(self.go('extract-pages',options={'pages':'2'}),1)
        self.assert_pages(self.go('organize',options={'pages':'2,1,1'}),3)
        split=self.go('split',options={'mode':'each'});self.assertEqual(split.suffix,'.zip')
        self.assert_pages(self.go('rotate',options={'angle':90}),2)
        self.assert_pages(self.go('crop',options={'left':10,'top':10,'right':10,'bottom':10}),2)
    def test_edit_tools(self):
        for slug,opts in [('watermark',{'text':'COPY'}),('page-numbers',{}),('edit',{'text':'EDITED'}),('sign',{'text':'Tester'}),('forms',{'field':'Full Name'})]:
            with self.subTest(slug=slug):
                self.assert_pages(self.go(slug,options=opts),2)
        redacted=self.go('redact',options={'text':'Secret Value'})
        with fitz.open(str(redacted)) as doc:self.assertNotIn('Secret Value',doc[1].get_text())
    def test_security_tools(self):
        secured=self.go('protect',options={'new_password':'demo1234'})
        with fitz.open(str(secured)) as doc:self.assertTrue(doc.needs_pass)
        unsealed=self.go('unlock',[secured],{'password':'demo1234'})
        with fitz.open(str(unsealed)) as doc:self.assertFalse(doc.needs_pass)
    def test_convert_image_text(self):
        self.assert_pages(self.go('jpg-to-pdf',[self.image]),1)
        self.assert_pages(self.go('scan',[self.image]),1)
        self.assertEqual(self.go('pdf-to-word').suffix,'.docx')
        self.assertEqual(self.go('pdf-to-powerpoint').suffix,'.pptx')
        self.assertEqual(self.go('pdf-to-excel').suffix,'.xlsx')
        self.assertEqual(self.go('pdf-to-markdown').suffix,'.md')
        self.assertEqual(self.go('summarize').suffix,'.txt')
        self.assertEqual(self.go('pdf-to-jpg').suffix,'.zip')
    def test_compare_compress_repair(self):
        self.assert_pages(self.go('compare',[self.pdf,self.second]),2)
        self.assert_pages(self.go('compress'),2)
        self.assert_pages(self.go('repair'),2)
    def test_office_to_pdf(self):
        doc=Document();doc.add_paragraph('A demo document');p=self.wd/'source.docx';doc.save(str(p))
        self.assert_pages(self.go('word-to-pdf',[p]),1)
        wb=Workbook();wb.active.append(['A','B']);p=self.wd/'sheet.xlsx';wb.save(str(p))
        self.assert_pages(self.go('excel-to-pdf',[p]),1)
        prs=Presentation();prs.slides.add_slide(prs.slide_layouts[0]);p=self.wd/'slides.pptx';prs.save(str(p))
        self.assert_pages(self.go('powerpoint-to-pdf',[p]),1)
    def test_html_to_pdf(self):
        self.assert_pages(self.go('html-to-pdf',[],{'html':'<h1>Hello from HTML</h1>'}),1)
    def test_pdfa(self):
        result=self.go('pdf-to-pdfa')
        self.assert_pages(result,2)
    def test_ocr(self):
        result=self.go('ocr',options={'dpi':100})
        self.assert_pages(result,2)
        with fitz.open(str(result)) as d:self.assertIn('Hello',d[0].get_text())

if __name__=='__main__':unittest.main()
